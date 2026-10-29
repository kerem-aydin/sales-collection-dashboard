"""Evaluate workbook formulas and compare them with independent raw-data totals."""
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re
import time
import traceback
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import formulas
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

ROOT = Path(__file__).resolve().parent
BOOK = ROOT / 'sales-collection-dashboard.xlsx'
REPORT = ROOT / 'docs' / 'verification.txt'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
CELL_KEY = re.compile(r"^'?\[[^\]]+\](.*?)'?!([A-Z]+[0-9]+)$")


def cents(value):
    return int((Decimal(str(value)) * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def independent(wb, month):
    """Read raw inputs without using formulas or importing the generator."""
    previous = datetime(int(month[:4]), int(month[5:]), 1)
    previous = f'{previous.year - (previous.month == 1):04d}-{(previous.month - 2) % 12 + 1:02d}'
    cities = [wb['Panel'].cell(r, 1).value for r in range(12, 22)]
    result = {city: {'sales': 0, 'payment': 0, 'previous': 0} for city in cities}
    customers = [tuple(wb['Müşteriler'].cell(r, c).value for c in range(1, 4))
                 for r in range(2, 102)]
    for date, invoice, customer, city, amount in wb['Satışlar'].iter_rows(
            min_row=2, max_row=2001, max_col=5, values_only=True):
        if date is None:
            continue
        assert isinstance(date, datetime), f'Invalid sales date: {invoice}'
        assert city in result, f'Sales province missing from panel: {city}'
        key = date.strftime('%Y-%m')
        if key == month:
            result[city]['sales'] += cents(amount)
        if key == previous:
            result[city]['previous'] += cents(amount)
    unmatched, outside, matches = [], [], []
    for r in range(2, 2002):
        date, desc, amount, balance = [wb['Banka Dökümü'].cell(r, c).value for c in range(1, 5)]
        if date is None or date.strftime('%Y-%m') != month:
            continue
        if amount <= 0:
            outside.append(r)
            continue
        # Match the sample aliases with case-insensitive substring search.
        found = [(name, city) for name, alias, city in customers
                 if alias and alias.upper() in (desc or '').upper()]
        if len(found) == 1 and found[0][0] and found[0][1]:
            name, city = found[0]
            assert city in result, f'Collection province missing from panel: {city}'
            result[city]['payment'] += cents(amount)
            matches.append((r, name, city))
        else:
            unmatched.append(r)
    for values in result.values():
        values['remaining'] = values['sales'] - values['payment']
        values['ratio'] = values['payment'] / values['sales'] if values['sales'] else 0
    return result, unmatched, outside, matches


def scalar(value):
    # Convert formulas ranges and NumPy arrays to Python scalars.
    v = value.value if hasattr(value, 'value') else value
    if hasattr(v, 'flatten'):
        v = v.flatten()[0]
    if hasattr(v, 'item'):
        v = v.item()
    return v


def cells(solution):
    values = {}
    for key, value in solution.items():
        if isinstance(key, str):
            match = CELL_KEY.match(key)
            if match:
                values[(match[1].upper(), match[2])] = scalar(value)
    return values


def cache_results(solution):
    """Cache actual engine results in OOXML while preserving formulas."""
    values = cells(solution)
    temp = ROOT / '.tmp' / 'cached.xlsx'
    temp.parent.mkdir(exist_ok=True)
    with ZipFile(BOOK) as source, ZipFile(temp, 'w', ZIP_DEFLATED) as target:
        workbook = ET.fromstring(source.read('xl/workbook.xml'))
        rels = ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        relmap = {r.attrib['Id']: r.attrib['Target'] for r in rels}
        sheetnames = {}
        for sheet in workbook.find('s:sheets', NS):
            rid = sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
            path = relmap[rid]
            path = path.lstrip('/') if path.startswith('/') else 'xl/' + path
            sheetnames[path] = sheet.attrib['name'].upper()
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename in sheetnames:
                tree = ET.fromstring(data)
                for cell in tree.findall('.//s:sheetData/s:row/s:c', NS):
                    if cell.find('s:f', NS) is None:
                        continue
                    key = (sheetnames[info.filename], cell.attrib['r'])
                    if key not in values:
                        continue
                    value = values[key]
                    if str(value).startswith(('#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#N/A', '#NUM!')):
                        raise AssertionError(f'Formula error: {key}: {value}')
                    for child in list(cell):
                        if child.tag in (f'{{{NS["s"]}}}v', f'{{{NS["s"]}}}is'):
                            cell.remove(child)
                    node = ET.SubElement(cell, f'{{{NS["s"]}}}v')
                    if isinstance(value, str):
                        cell.set('t', 'str')
                        node.text = value
                    elif isinstance(value, bool):
                        cell.set('t', 'b')
                        node.text = '1' if value else '0'
                    else:
                        cell.attrib.pop('t', None)
                        node.text = str(float(value))
                data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            target.writestr(info, data)
    temp.replace(BOOK)


def assert_money(actual, expected, label):
    assert cents(actual) == expected, f'{label}: formula={actual}, independent={expected / 100}'


def boundary_checks(model, month_key, wb, summary):
    """Exercise boundary inputs without changing the delivered workbook."""
    prefix = month_key.split(']')[0] + ']'
    def key(sheet, address):
        return prefix + sheet.upper() + "'!" + address

    def evaluate(inputs, outputs):
        return cells(model.calculate(inputs={key(s, a): v for (s, a), v in inputs.items()},
                                     outputs=[key(s, a) for s, a in outputs]))
    default = {('Panel', 'B4'): '2026-10'}
    expected = summary['2026-10']
    # Exercise the last transaction row and the last customer slot.
    inputs = {**default,
        ('Müşteriler', 'A101'): 'Kapasite Örnek Müşteri',
        ('Müşteriler', 'B101'): 'KAPASITE ORNEK 2000',
        ('Müşteriler', 'C101'): 'Ankara',
        ('Banka Dökümü', 'A2001'): to_excel(datetime(2026, 10, 29)),
        ('Banka Dökümü', 'B2001'): 'FAST GELEN - KAPASITE ORNEK 2000',
        ('Banka Dökümü', 'C2001'): 123.45,
        ('Satışlar', 'A2001'): to_excel(datetime(2026, 10, 29)),
        ('Satışlar', 'D2001'): 'Ankara', ('Satışlar', 'E2001'): 250.00}
    a = evaluate(inputs, [('Panel', 'A7'), ('Panel', 'C7'), ('Panel', 'N4'),
                          ('Banka Dökümü', 'H2001')])
    assert_money(a['PANEL', 'A7'], expected['sales'] + 25000, '2.000 satış kapasitesi')
    assert_money(a['PANEL', 'C7'], expected['payment'] + 12345, '2.000 banka kapasitesi')
    assert a['BANKA DÖKÜMÜ', 'H2001'] == 'Eşleşti'
    assert int(a['PANEL', 'N4']) == len(expected['unmatched'])
    # Ambiguous descriptions must not select the first customer.
    row = next(r for r in range(2, 112) if wb['Banka Dökümü'][f'C{r}'].value > 0)
    desc = wb['Müşteriler']['B2'].value + ' / ' + wb['Müşteriler']['B3'].value
    a = evaluate({**default, ('Banka Dökümü', f'B{row}'): desc},
                 [('Banka Dökümü', f'F{row}'), ('Banka Dökümü', f'H{row}')])
    assert a['BANKA DÖKÜMÜ', f'F{row}'] == ''
    assert a['BANKA DÖKÜMÜ', f'H{row}'] == 'Eşleşmedi'
    # Adding a customer must resolve an existing unmatched payment.
    row = expected['unmatched'][0]
    inputs = {**default, ('Müşteriler', 'A101'): 'Yeni Örnek Müşteri',
              ('Müşteriler', 'B101'): wb['Banka Dökümü'][f'B{row}'].value,
              ('Müşteriler', 'C101'): 'Ankara'}
    a = evaluate(inputs, [('Panel', 'C7'), ('Panel', 'N4')])
    assert int(a['PANEL', 'N4']) == len(expected['unmatched']) - 1
    assert_money(a['PANEL', 'C7'], expected['payment'] + cents(wb['Banka Dökümü'][f'C{row}'].value),
                 'Yeni müşteriyle eşleştirme')
    # Empty months have zero totals and a blank change percentage.
    a = evaluate({('Panel', 'B4'): '2026-01'}, [('Panel', c) for c in ['A7', 'C7', 'E7', 'H7', 'N4']])
    assert all(a['PANEL', c] == 0 for c in ['A7', 'C7', 'E7', 'N4'])
    assert a['PANEL', 'H7'] == ''


def run():
    REPORT.parent.mkdir(exist_ok=True)
    start = time.monotonic()
    lines = ['SALES AND COLLECTIONS / VERIFICATION', f'Date: {datetime.now():%Y-%m-%d %H:%M:%S}',
             f'Workbook: {BOOK.name}', f'Calculation engine: formulas {formulas.__version__}',
             'Comparison: evaluated workbook formulas / independent raw-input calculation (integer cents).', '']
    with ZipFile(BOOK) as archive:
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')):
                xml = archive.read(name).decode('utf-8').lower()
                assert not re.search(r'spire|evaluation warning|evaluation version|evaluation worksheet', xml), name
    wb = load_workbook(BOOK, data_only=False)
    assert wb.sheetnames == ['Panel', 'Banka Dökümü', 'Satışlar', 'Müşteriler', 'Eşleştirme', 'Nasıl kullanılır']
    assert wb.active.title == 'Panel'
    assert not wb.vba_archive
    assert len(wb['Panel']._charts) == 1
    assert wb['Panel']._charts[0].type == 'col'
    assert wb['Panel'].sheet_view.showGridLines is False
    assert wb['Panel']['B4'].value == '2026-10'
    assert any(str(v.sqref) == 'B4' and v.type == 'list' for v in wb['Panel'].data_validations.dataValidation)
    for sheet, columns in [('Banka Dökümü', 'EFGHI'), ('Satışlar', 'F')]:
        for col in columns:
            assert all(wb[sheet][f'{col}{r}'].data_type == 'f' for r in range(2, 2002))
    forbidden = re.compile(r'\b(XLOOKUP|LET|FILTER|SORT|UNIQUE|SEQUENCE)\s*\(', re.I)
    count = 0
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    count += 1
                    assert not forbidden.search(cell.value), f'Unsupported formula: {ws.title}!{cell.coordinate}'
    assert count == 22065, f'Unexpected formula count: {count}'
    lines += [f'Structure: 6 sheets, 1 native column chart, month validation list, {count} formulas.',
              'Capacity: 2,000 sales rows, 2,000 bank rows and 100 customer slots.',
              'No VBA, dynamic-array functions or evaluation markers.', '']
    print('Loading formula model...', flush=True)
    model = formulas.ExcelModel().loads(str(BOOK)).finish()
    month_key = next(k for k in model.cells if k.endswith("PANEL'!B4"))
    all_results = {}
    summary = {}
    for month in ['2026-09', '2026-10']:
        print(f'{month}: calculating...', flush=True)
        solution = model.calculate(inputs={month_key: month})
        actual = cells(solution)
        def get(sheet, cell):
            return actual[(sheet.upper(), cell)]
        expected, unmatched, outside, matches = independent(wb, month)
        checks = 0
        for r, (city, values) in enumerate(expected.items(), 12):
            for col, key in [('B', 'sales'), ('C', 'payment'), ('D', 'remaining'), ('F', 'previous')]:
                assert_money(get('Panel', f'{col}{r}'), values[key], f'{month}/{city}/{key}')
                checks += 1
            assert abs(float(get('Panel', f'E{r}')) - values['ratio']) < 1e-12
            checks += 1
        sales = sum(v['sales'] for v in expected.values())
        paid = sum(v['payment'] for v in expected.values())
        previous = sum(v['previous'] for v in expected.values())
        for address, exp in [('A7', sales), ('C7', paid), ('E7', sales - paid),
                             ('B22', sales), ('C22', paid), ('D22', sales - paid), ('F22', previous)]:
            assert_money(get('Panel', address), exp, f'{month}/{address}')
            checks += 1
        ratio = paid / sales if sales else 0
        assert abs(float(get('Panel', 'E22')) - ratio) < 1e-12
        if previous:
            assert abs(float(get('Panel', 'H7')) - (sales / previous - 1)) < 1e-12
        else:
            assert get('Panel', 'H7') == ''
        assert int(get('Panel', 'N4')) == len(unmatched)
        assert f'{len(unmatched)} ödeme' in get('Panel', 'A24')
        checks += 4
        for r, name, city in matches:
            assert get('Banka Dökümü', f'F{r}') == name
            assert get('Banka Dökümü', f'G{r}') == city
            assert get('Banka Dökümü', f'H{r}') == 'Eşleşti'
        for r in unmatched:
            assert get('Banka Dökümü', f'H{r}') == 'Eşleşmedi'
        for r in outside:
            assert get('Banka Dökümü', f'H{r}') == 'Satış dışı'
        for offset, source_row in enumerate(unmatched, 5):
            assert get('Eşleştirme', f'B{offset}') == wb['Banka Dökümü'][f'B{source_row}'].value
            assert_money(get('Eşleştirme', f'C{offset}'), cents(wb['Banka Dökümü'][f'C{source_row}'].value), 'Eşleştirme tutarı')
        assert get('Eşleştirme', f'B{5 + len(unmatched)}') == ''
        assert get('Satışlar', 'F2001') == ''
        assert get('Banka Dökümü', 'H2001') == ''
        errors = [(k, str(v)) for k, v in solution.items() if re.search(
                  r'\[\[.?\#(REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!)', str(v))]
        assert not errors, f'Formula errors: {errors[:5]}'
        lines += [f'{month}: PASS — {checks} panel comparisons across all 10 cities.',
                  f'  Sales: {sales / 100:.2f} TRY | Collections: {paid / 100:.2f} TRY | Remaining: {(sales - paid) / 100:.2f} TRY',
                  f'  Previous-month sales: {previous / 100:.2f} TRY | Unmatched: {len(unmatched)} | Non-sales: {len(outside)}',
                  f'  {len(matches)} matched transactions, {len(unmatched)} review rows, empty capacity and error scan: PASS', '']
        for city, v in expected.items():
            lines.append(f'  {city}: sales {v["sales"] / 100:.2f} / collections {v["payment"] / 100:.2f} / remaining {v["remaining"] / 100:.2f} TRY')
        lines.append('')
        summary[month] = {'cities': expected, 'unmatched': unmatched, 'outside': outside,
                          'sales': sales, 'payment': paid, 'previous': previous}
        all_results[month] = solution
        print(f'{month}: PASS ({checks} panel comparisons)', flush=True)
    print('Checking capacity and matching boundaries...', flush=True)
    boundary_checks(model, month_key, wb, summary)
    lines += ['Additional boundary checks: PASS',
              '  Last sales/bank row and last customer slot contribute to totals.',
              '  Ambiguous aliases rejected; new customer resolves unmatched payment.',
              '  Empty month returns zero totals and a blank change percentage.', '']
    print('Caching verified October results...', flush=True)
    cache_results(all_results['2026-10'])
    cached = load_workbook(BOOK, data_only=True)
    for coordinate in ['A7', 'C7', 'E7', 'H7', 'B12', 'N4']:
        assert cached['Panel'][coordinate].value == cells(all_results['2026-10'])[('PANEL', coordinate)]
    formulas_after = load_workbook(BOOK, data_only=False)
    expected_cache = cells(all_results['2026-10'])
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    assert formulas_after[sheet.title][cell.coordinate].value == cell.value
                    expected_value = expected_cache[sheet.title.upper(), cell.coordinate]
                    actual_value = cached[sheet.title][cell.coordinate].value
                    if isinstance(actual_value, datetime):
                        actual_value = to_excel(actual_value, cached.epoch)
                    if expected_value == '':
                        assert actual_value in ('', None)
                    else:
                        assert actual_value == expected_value, (sheet.title, cell.coordinate, actual_value, expected_value)
    lines += ['Opening cache: verified October engine results written; all 22,065 formulas preserved and cached.',
              f'XLSX SHA256: {hashlib.sha256(BOOK.read_bytes()).hexdigest()}',
              f'Duration: {time.monotonic() - start:.1f} seconds',
              'This check does not cover visual or interactive testing in Microsoft Excel.',
              '', 'RESULT: PASS']
    (ROOT / '.tmp' / 'verification-data.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    REPORT.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print('RESULT: PASS — 122 panel comparisons plus matching and boundary checks', flush=True)


if __name__ == '__main__':
    try:
        run()
    except Exception:
        error = traceback.format_exc()
        REPORT.parent.mkdir(exist_ok=True)
        REPORT.write_text('RESULT: FAIL\n\n' + error, encoding='utf-8')
        raise
