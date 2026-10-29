"""Generate a Turkish sales and collections workbook with fictional data."""
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'sales-collection-dashboard.xlsx'
LIMIT = 2001  # Header plus 2,000 data rows.
CUSTOMER_LIMIT = 101
CITIES = ['İstanbul', 'Ankara', 'İzmir', 'Bursa', 'Antalya', 'Konya',
          'Kocaeli', 'Sakarya', 'Kayseri', 'Gaziantep']
NAVY, INK, MUTED, GRAY, LINE = '1F3A5F', '233247', '64748B', 'F2F4F7', 'DFE5EB'
MONEY = '#,##0.00 "₺"'
PERCENT = '0.0%'
THIN = Side(style='thin', color=LINE)


def sample_data():
    rng = random.Random(20261029)
    names = ['Mavidal Gıda', 'Düşova Market', 'Pusuyel Tedarik', 'Çınariz Dağıtım',
             'Ayvenk Ticaret', 'Kumsera Gıda', 'Yelmino Market', 'Morçatı Tedarik',
             'Gölvera Dağıtım', 'Narvadi Ticaret', 'Çiyrota Gıda', 'Bulutiz Market',
             'Kervola Tedarik', 'Ayserin Dağıtım', 'Yapriva Ticaret', 'Denizova Gıda',
             'Koralin Market', 'Zerova Tedarik', 'Ardılyel Dağıtım', 'Günsera Ticaret',
             'Tanyuva Gıda', 'Erguven Market', 'Kumvadi Tedarik', 'Lalrota Dağıtım',
             'Çisenta Ticaret']
    trans = str.maketrans('çğıöşüÇĞİÖŞÜ', 'cgiosuCGIOSU')
    customers = [(f'{n} Örnek İşletme', f'KURGUSAL {n.translate(trans).upper()}',
                  CITIES[i % 10]) for i, n in enumerate(names)]
    sales, bank = [], []
    for month in (9, 10):
        indices = list(range(25)) * 2 + [rng.randrange(25) for _ in range(20)]
        month_sales = []
        for i, idx in enumerate(indices, 1):
            name, _, city = customers[idx]
            cents = rng.randint(325000, 1950000)
            month_sales.append((datetime(2026, month, rng.randint(1, 28)),
                                f'TSL-2026{month:02d}-{i:03d}', name, city, cents / 100))
        sales.extend(sorted(month_sales))
        totals = defaultdict(int)
        for _, _, name, _, value in month_sales:
            totals[name] += round(value * 100)
        paid_indices = list(range(25)) + [rng.randrange(25) for _ in range(24)]
        counts = Counter(paid_indices)
        month_bank = []
        for i, idx in enumerate(paid_indices, 1):
            name, alias, _ = customers[idx]
            # Collect roughly 62-84% of monthly sales per customer.
            cents = round(totals[name] * (0.62 + (idx % 12) * .02) / counts[idx])
            description = [f'FAST GELEN - {alias} / FATURA ODEMESI',
                           f'EFT - GONDEREN: {alias} - CARI HESAP',
                           f'HAVALE {alias} / MAL BEDELI'][i % 3]
            month_bank.append((datetime(2026, month, rng.randint(2, 28)), description, cents / 100))
        for i, alias in enumerate(['BELIRSIZ GONDEREN 710', 'TANIMSIZ ISLETME 820',
                                   'ACIKLAMASIZ TRANSFER 930']):
            month_bank.append((datetime(2026, month, 8 + i * 6), f'FAST GELEN - {alias}',
                               (525000 + i * 173050 + month * 1000) / 100))
        month_bank.extend([(datetime(2026, month, 5), 'EFT ISLEM UCRETI / BANKA MASRAFI', -28.50),
                           (datetime(2026, month, 15), 'MAAS ODEMESI - PERSONEL TOPLU', -32750),
                           (datetime(2026, month, 20), 'HAVALE - ISYERI KIRA ODEMESI', -18500)])
        bank.extend(sorted(month_bank))
    balance = 24500000
    bank_rows = []
    for date, desc, amount in sorted(bank):
        balance += round(amount * 100)
        bank_rows.append((date, desc, amount, balance / 100))
    return customers, sorted(sales), bank_rows


def base_sheet(ws, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = 'B2'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.sheet_properties.outlinePr.summaryRight = False


def header(ws, labels):
    ws.append(labels)
    ws.row_dimensions[1].height = 30
    for c in ws[1]:
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        c.alignment = Alignment(vertical='center', wrap_text=True)


def data_style(ws, first, last, count, input_cols):
    for row in ws.iter_rows(min_row=first, max_row=last, max_col=count):
        for c in row:
            c.font = Font(name='Calibri', size=11, color=INK)
            c.alignment = Alignment(vertical='center')
            if c.column > input_cols:
                c.fill = PatternFill('solid', fgColor=GRAY)
        ws.row_dimensions[row[0].row].height = 21


def add_table(ws, name, ref):
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True,
                                        showFirstColumn=False, showLastColumn=False)
    ws.add_table(tab)


def named(wb, name, ref):
    wb.defined_names.add(DefinedName(name, attr_text=ref))


def list_validation(ws, formula, ref):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = 'Listeden seçim yapın'
    dv.error = 'Tanımlı değerlerden birini seçin; yeni kayıt için ayarlar sayfasını güncelleyin.'
    dv.errorStyle = 'stop'
    ws.add_data_validation(dv)
    dv.add(ref)


def merged(ws, ref, value, size=11, bold=False, color=INK, fill=None, align='left'):
    ws.merge_cells(ref)
    cell = ws[ref.split(':')[0]]
    cell.value = value
    cell.font = Font(name='Calibri', size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical='center')
    if fill:
        for row in ws[ref]:
            for c in row:
                c.fill = PatternFill('solid', fgColor=fill)
    return cell


def style_card_boundaries(panel):
    for col in ['B', 'D', 'F']:
        for row in range(6, 9):
            panel[f'{col}{row}'].border = Border(right=Side(style='medium', color='FFFFFF'))


def create_workbook():
    customers, sales, bank = sample_data()
    wb = Workbook()
    panel = wb.active
    panel.title = 'Panel'
    bs = wb.create_sheet('Banka Dökümü')
    ss = wb.create_sheet('Satışlar')
    cs = wb.create_sheet('Müşteriler')
    es = wb.create_sheet('Eşleştirme')
    help_sheet = wb.create_sheet('Nasıl kullanılır')
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True,
                                    calcMode='auto')
    wb.properties.title = 'İl Bazlı Satış ve Tahsilat • TASLAK'
    wb.properties.subject = 'Tamamı kurgusal örnek veridir; gerçek müşteri verisi içermez.'
    wb.properties.creator = 'Satış ve Tahsilat Taslağı'

    base_sheet(cs, {'A': 38, 'B': 34, 'C': 18, 'E': 3, 'F': 18, 'H': 16})
    header(cs, ['Müşteri', 'Aranacak ad', 'İl'])
    for row in customers:
        cs.append(row)
    data_style(cs, 2, CUSTOMER_LIMIT, 3, 3)
    cs['F1'] = 'İl listesi'
    cs['H1'] = 'Ay listesi'
    for i, city in enumerate(CITIES, 2):
        cs.cell(i, 6, city)
    for i in range(1, 13):
        cs.cell(i + 1, 8, f'2026-{i:02d}')
        cs.cell(i + 1, 8).number_format = '@'
    named(wb, 'Iller', "'Müşteriler'!$F$2:$F$11")
    named(wb, 'Aylar', "'Müşteriler'!$H$2:$H$13")
    named(wb, 'Musteriler', "'Müşteriler'!$A$2:$A$101")
    list_validation(cs, '=Iller', 'C2:C101')
    cs['B1'].comment = Comment('En fazla 100 müşteri. Aranacak ad dolu, benzersiz ve ayırt edici olmalı. '
                               'SEARCH büyük/küçük harf duyarsızdır. *, ? ve ~ kullanmayın. '
                               'Birden fazla ad bulunursa ödeme eşleşmez.', 'User')
    add_table(cs, 'MusteriAyar', 'A1:C101')

    base_sheet(ss, {'A': 14, 'B': 25, 'C': 38, 'D': 18, 'E': 20, 'F': 14})
    header(ss, ['Tarih', 'Fatura no', 'Müşteri', 'İl', 'Tutar', 'Ay'])
    for row in sales:
        ss.append(row)
    data_style(ss, 2, LIMIT, 6, 5)
    for r in range(2, LIMIT + 1):
        ss.cell(r, 6, f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))')
        ss.cell(r, 1).number_format = 'dd.mm.yyyy'
        ss.cell(r, 5).number_format = MONEY
    list_validation(ss, '=Iller', 'D2:D2001')
    list_validation(ss, '=Musteriler', 'C2:C2001')
    add_table(ss, 'SatisVerisi', 'A1:F2001')
    ss['F1'].comment = Comment('Gri sütun formüldür; A:E alanına yapıştırın. 2.000 veri satırı hazırdır.', 'User')

    base_sheet(bs, {'A': 14, 'B': 64, 'C': 20, 'D': 20, 'E': 14, 'F': 38, 'G': 18, 'H': 18, 'I': 15})
    header(bs, ['Tarih', 'Açıklama', 'Tutar', 'Bakiye', 'Ay', 'Bulunan müşteri', 'İl', 'Durum', 'İnceleme sırası'])
    for row in bank:
        bs.append(row)
    data_style(bs, 2, LIMIT, 9, 4)
    keys = "'Müşteriler'!$B$2:$B$101"
    for r in range(2, LIMIT + 1):
        hits = f'({keys}<>"")*ISNUMBER(SEARCH({keys},B{r}))'
        bs.cell(r, 5, f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))')
        bs.cell(r, 6, f'=IF(OR(A{r}="",B{r}="",C{r}<=0),"",'
                f'IF(SUMPRODUCT({hits})=1,IFERROR(INDEX(\'Müşteriler\'!$A$2:$A$101,'
                f'SUMPRODUCT({hits},ROW({keys})-ROW(\'Müşteriler\'!$B$2)+1)),""),""))')
        bs.cell(r, 7, f'=IF(F{r}="","",IFERROR(INDEX(\'Müşteriler\'!$C$2:$C$101,'
                f'MATCH(F{r},\'Müşteriler\'!$A$2:$A$101,0)),""))')
        bs.cell(r, 8, f'=IF(A{r}="","",IF(C{r}<=0,"Satış dışı",'
                f'IF(OR(F{r}="",G{r}=""),"Eşleşmedi","Eşleşti")))')
        # A running counter locates each unmatched transaction.
        # Linear dependencies avoid expanding COUNTIFS ranges.
        before = '0' if r == 2 else f'I{r - 1}'
        bs.cell(r, 9, f'={before}+IF(AND(H{r}="Eşleşmedi",E{r}=\'Panel\'!$B$4),1,0)')
        bs.cell(r, 1).number_format = 'dd.mm.yyyy'
        for col in (3, 4):
            bs.cell(r, col).number_format = MONEY
    bs.column_dimensions['I'].hidden = True
    bs.conditional_formatting.add('H2:H2001', FormulaRule(formula=['H2="Eşleşmedi"'],
        font=Font(color=NAVY, bold=True), fill=PatternFill('solid', fgColor='E5EBF2')))
    bs['H1'].comment = Comment('Yalnız tek müşteriyle eşleşen pozitif girişler tahsilattır. '
        'Sıfır/negatif tutarlar Satış dışı; adsız veya çoklu eşleşen girişler Eşleşmedi olur.', 'User')
    add_table(bs, 'BankaVerisi', 'A1:H2001')

    base_sheet(panel, {'A': 18, 'B': 17, 'C': 17, 'D': 18, 'E': 16, 'F': 19,
                       'G': 3, 'H': 12, 'I': 12, 'J': 12, 'K': 12, 'L': 12})
    panel.freeze_panes = 'B12'
    panel.sheet_view.zoomScale = 85
    panel.sheet_view.zoomScaleNormal = 85
    panel.sheet_view.topLeftCell = 'A1'
    panel.sheet_view.selection[0].activeCell = 'B4'
    panel.sheet_view.selection[0].sqref = 'B4'
    heights = {1: 25, 2: 17, 3: 6, 4: 23, 5: 6, 6: 16, 7: 30, 8: 16,
               9: 7, 10: 18, 11: 28, **{r: 18 for r in range(12, 22)},
               22: 22, 23: 6, 24: 25, 25: 16, 26: 16}
    for r, h in heights.items():
        panel.row_dimensions[r].height = h
    merged(panel, 'A1:L1', 'SATIŞ & TAHSİLAT', 20, True, NAVY)
    merged(panel, 'A2:L2', 'İl bazında aylık görünüm   /   TASLAK · Tamamı kurgusal örnek veridir', 11, color=MUTED)
    panel['A4'] = 'GÖRÜNTÜLENEN AY'
    panel['A4'].font = Font(name='Calibri', size=10, bold=True, color=MUTED)
    panel['B4'] = '2026-10'
    panel['B4'].font = Font(name='Calibri', size=12, bold=True, color=NAVY)
    panel['B4'].fill = PatternFill('solid', fgColor='E5EBF2')
    panel['B4'].alignment = Alignment(horizontal='center', vertical='center')
    panel['B4'].border = Border(bottom=Side(style='thin', color=NAVY))
    list_validation(panel, '=Aylar', 'B4')
    merged(panel, 'C4:F4', 'Ayı listeden seçin · Tüm özetler ve grafik güncellenir', 10, color=MUTED)
    merged(panel, 'H4:L4', 'VERİ KAPSAMI  /  EYLÜL + EKİM 2026', 10, True, MUTED)
    panel['N2'] = '=TEXT(DATE(LEFT(B4,4),RIGHT(B4,2)-1,1),"yyyy-mm")'
    panel['N3'] = '=SUM(F12:F21)'
    panel['N4'] = '=COUNTIFS(\'Banka Dökümü\'!$E$2:$E$2001,$B$4,\'Banka Dökümü\'!$H$2:$H$2001,"Eşleşmedi")'
    panel.column_dimensions['N'].hidden = True
    cards = [('A', 'B', 'Toplam satış', '=SUM(B12:B21)', 'Seçilen ayın faturaları', MONEY),
             ('C', 'D', 'Gelen ödeme', '=SUM(C12:C21)', 'Eşleşen pozitif banka girişleri', MONEY),
             ('E', 'F', 'Kalan alacak', '=A7-C7', 'Seçilen ay: satış − tahsilat', MONEY),
             ('H', 'L', 'Geçen aya göre satış farkı %', '=IF($N$3=0,"",A7/$N$3-1)',
              '=IF($N$3=0,"Önceki ay satış yok / oran hesaplanmadı","Önceki ay satışına göre değişim")', PERCENT)]
    for start, end, title, formula, desc, fmt in cards:
        merged(panel, f'{start}6:{end}6', title, 11, True, NAVY, GRAY)
        merged(panel, f'{start}7:{end}7', formula, 22, True, NAVY, GRAY).number_format = fmt
        merged(panel, f'{start}8:{end}8', desc, 10, color=MUTED, fill=GRAY)
    style_card_boundaries(panel)
    merged(panel, 'A10:F10', 'İLLERE GÖRE AYLIK ÖZET', 11, True, NAVY)
    merged(panel, 'H10:L10', 'İL BAZLI SATIŞ  /  ₺', 11, True, NAVY)
    for col, label in enumerate(['İl', 'Satış', 'Ödeme', 'Kalan alacak', 'Tahsilat oranı', 'Geçen ay satış'], 1):
        c = panel.cell(11, col, label)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.font = Font(name='Calibri', size=11, color='FFFFFF', bold=True)
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for r, city in enumerate(CITIES, 12):
        panel.cell(r, 1, city)
        panel.cell(r, 2, f'=SUMIFS(\'Satışlar\'!$E$2:$E$2001,\'Satışlar\'!$F$2:$F$2001,$B$4,\'Satışlar\'!$D$2:$D$2001,A{r})')
        panel.cell(r, 3, f'=SUMIFS(\'Banka Dökümü\'!$C$2:$C$2001,\'Banka Dökümü\'!$E$2:$E$2001,$B$4,'
                           f'\'Banka Dökümü\'!$G$2:$G$2001,A{r},\'Banka Dökümü\'!$H$2:$H$2001,"Eşleşti")')
        panel.cell(r, 4, f'=B{r}-C{r}')
        panel.cell(r, 5, f'=IFERROR(C{r}/B{r},0)')
        panel.cell(r, 6, f'=SUMIFS(\'Satışlar\'!$E$2:$E$2001,\'Satışlar\'!$F$2:$F$2001,$N$2,\'Satışlar\'!$D$2:$D$2001,A{r})')
        for c in panel[r][:6]:
            c.font = Font(name='Calibri', size=11, color=INK)
            c.alignment = Alignment(vertical='center', horizontal='left' if c.column == 1 else 'right')
            c.fill = PatternFill('solid', fgColor='FFFFFF' if r % 2 == 0 else GRAY)
            c.border = Border(bottom=Side(style='hair', color=LINE))
            if c.column > 1:
                c.number_format = PERCENT if c.column == 5 else MONEY
    panel.conditional_formatting.add('B12:B21', DataBarRule(start_type='num', start_value=0,
                                  end_type='max', color=NAVY, showValue=True))
    panel['A22'] = 'TOPLAM'
    for col in (2, 3, 4, 6):
        letter = chr(64 + col)
        panel.cell(22, col, f'=SUM({letter}12:{letter}21)').number_format = MONEY
    panel['E22'] = '=IFERROR(C22/B22,0)'
    panel['E22'].number_format = PERCENT
    for c in panel[22][:6]:
        c.fill = PatternFill('solid', fgColor=GRAY)
        c.font = Font(name='Calibri', size=11, bold=True, color=NAVY)
        c.border = Border(top=THIN)
    merged(panel, 'A24:L24', '=IF($N$4=0,"Eşleşmeyen ödeme yok → Eşleştirme sayfası",'
           '"Eşleşmeyen "&$N$4&" ödeme var → Eşleştirme sayfası")', 12, True, NAVY, GRAY)
    panel['A24'].hyperlink = "#'Eşleştirme'!A1"
    merged(panel, 'A25:L25', 'Sabit il sırası · Satış çubukları büyüklüğü gösterir · Eşleşmeyen girişler tahsilata dahil değildir.', 10, color=MUTED)
    merged(panel, 'A26:L26', 'Alacak yalnız seçilen ayın farkıdır; önceki dönem devri, vade ve fatura bazlı kapama bu taslakta yoktur.', 10, color=MUTED)
    chart = BarChart()
    chart.type = 'col'
    chart.style = 13
    chart.title = None
    chart.y_axis.title = None
    chart.y_axis.numFmt = '#,##0," bin ₺"'
    chart.y_axis.scaling.min = 0
    chart.x_axis.title = None
    chart.legend = None
    chart.width, chart.height = 11.7, 8.0
    chart.gapWidth = 65
    chart.add_data(Reference(panel, min_col=2, min_row=11, max_row=21), titles_from_data=True)
    chart.set_categories(Reference(panel, min_col=1, min_row=12, max_row=21))
    chart.series[0].graphicalProperties = GraphicalProperties(solidFill=NAVY, ln=GraphicalProperties().ln)
    chart.series[0].graphicalProperties.line.noFill = True
    panel.add_chart(chart, 'H11')
    panel.print_area = 'A1:L26'
    panel.page_setup.orientation = 'landscape'
    panel.page_setup.paperSize = panel.PAPERSIZE_A4
    panel.page_setup.fitToWidth = panel.page_setup.fitToHeight = 1
    panel.print_options.horizontalCentered = True
    panel.sheet_properties.tabColor = NAVY

    base_sheet(es, {'A': 14, 'B': 64, 'C': 20, 'D': 56, 'E': 10})
    header(es, ['Tarih', 'Açıklama', 'Tutar', 'Yapılacak', 'Kaynak satır'])
    es.freeze_panes = 'A5'
    merged(es, 'A2:D2', '=\'Panel\'!$B$4&"  /  "&\'Panel\'!$N$4&" eşleşmeyen ödeme"', 12, True, NAVY)
    merged(es, 'A3:D3', 'Aranacak adı Müşteriler sayfasına ekleyin/düzeltin; liste kendiliğinden güncellenir.', 11, color=MUTED)
    es.row_dimensions[2].height = 25
    es.row_dimensions[3].height = 22
    for r in range(5, LIMIT + 4):
        idx = r - 4
        es.cell(r, 5, f'=IFERROR(MATCH({idx},\'Banka Dökümü\'!$I$2:$I$2001,0),"")')
        for col, source in [(1, 'A'), (2, 'B'), (3, 'C')]:
            es.cell(r, col, f'=IF($E{r}="","",INDEX(\'Banka Dökümü\'!${source}$2:${source}$2001,$E{r}))')
        es.cell(r, 4, f'=IF(E{r}="","","Müşteri adı / aranacak ad / il bilgisini kontrol edin")')
        es.cell(r, 1).number_format = 'dd.mm.yyyy'
        es.cell(r, 3).number_format = MONEY
    data_style(es, 5, LIMIT + 3, 5, 0)
    es.column_dimensions['E'].hidden = True
    es.auto_filter.ref = f'A4:D{LIMIT + 3}'
    for col, label in enumerate(['Tarih', 'Açıklama', 'Tutar', 'Yapılacak'], 1):
        es.cell(4, col, label).font = Font(name='Calibri', bold=True, color=NAVY)
    es['A1'].hyperlink = "#'Panel'!A1"

    base_sheet(help_sheet, {'A': 6, 'B': 112})
    help_sheet.freeze_panes = 'A3'
    merged(help_sheet, 'A1:B1', 'NASIL KULLANILIR  /  5 ADIM', 18, True, NAVY)
    help_sheet.row_dimensions[1].height = 34
    steps = [
        'Banka Dökümü sayfasında A:D alanına gerçek tarih, açıklama, tutar ve bakiye yapıştırın. Girişler pozitif, çıkışlar negatif olmalı; E:I formüllerini silmeyin.',
        'Satışlar sayfasında A:E alanına tarih, fatura no, müşteri, il ve tutar girin. F sütunu ayı hesaplar. Eski veriyi değiştiriyorsanız kalan eski satırları da temizleyin.',
        'Müşteriler sayfasına müşteri, bankada aranacak ayırt edici kısa ad ve il ekleyin. Boş, yinelenen veya birden çok müşteriyi çağrıştıran ad kullanmayın (en fazla 100 müşteri).',
        'Panel sayfasında B4 hücresinden ayı seçin. Satış, tahsilat, ay içi alacak farkı, önceki ay karşılaştırması ve grafik birlikte güncellenir.',
        'Paneldeki uyarıya tıklayın. Eşleştirme listesindeki girişler için müşteri/aranacak ad/il bilgisini düzeltin; tek eşleşme sağlanınca ödeme tahsilata katılır.'
    ]
    for r, text in enumerate(steps, 3):
        help_sheet.cell(r, 1, str(r - 2).zfill(2)).font = Font(name='Calibri', size=16, bold=True, color=NAVY)
        c = help_sheet.cell(r, 2, text)
        c.font = Font(name='Calibri', size=12, color=INK)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        help_sheet.row_dimensions[r].height = 58
    merged(help_sheet, 'A10:B10', 'TASLAK • Örnek işletmeler ve işlemler kurgusaldır. Gerçek müşteri verisi yoktur.', 11, True, NAVY)
    merged(help_sheet, 'A12:B12', 'Kapasite: 2.000 satış + 2.000 banka satırı / 100 müşteri. Gri alanlar hesaplanır; formüllere yapıştırmayın.', 11, color=MUTED)
    merged(help_sheet, 'A13:B13', 'Alacak = seçilen ay satışı − o ayın eşleşen tahsilatı. Açılış bakiyesi ve fatura/vade takibi içermez.', 11, color=MUTED)
    merged(help_sheet, 'A14:B14', 'Yeni ay/yıl için Müşteriler!H2:H13 ay listesini güncelleyin. İller bu taslakta sabit 10 ildir.', 11, color=MUTED)
    wb.active = 0
    wb.save(OUTPUT)
    print(f'GENERATED: {OUTPUT.name} | {len(customers)} customers, {len(sales)} sales, {len(bank)} bank transactions')
    return OUTPUT


if __name__ == '__main__':
    create_workbook()
